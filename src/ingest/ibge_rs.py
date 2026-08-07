"""Ingestao das bases municipais do RS (IBGE) — Camada 1, dominio municipal.

Por que este modulo nao usa `base.Ingestor`: aquele contrato termina em
`validate_series`, isto e, no esquema (timestamp, value, quality_flag). As
bases aqui nao sao serie temporal — sao um retrato transversal por municipio
(497 linhas) e uma malha geografica. Forcar o esquema de serie exigiria
inventar um `timestamp` por linha, que e exatamente o tipo de metadado falso
que a proveniencia existe para impedir.

O que E reaproveitado, porque e o que importa: a MESMA disciplina de
`base.Ingestor` — payload bruto imutavel em `data/raw/<source>/<ts>/` com
`provenance.json` (url, sha256, fetched_at, rows), cache por janela, respeito
a `CLIMATE_OFFLINE=1`, e so entao o parse para `data/interim/`. Um parser que
quebrar amanha reprocessa o bruto de hoje sem rede.

Fontes
------
`ibge_munic_rs`
    Pesquisa de Informacoes Basicas Municipais (MUNIC) 2024, suplemento
    "Evento Climatico Rio Grande do Sul": 497 municipios x 100 variaveis
    levantadas junto as prefeituras sobre o evento iniciado em 26/04/2024.
    E a unica base publica que responde, por municipio, TRES perguntas que
    esta central precisa e nao tinha: quais perigos hidricos ocorreram, qual
    foi o impacto, e — decisivo — se havia plano de contingencia, se ele foi
    executado e, quando nao foi, POR QUE nao (falta de recurso financeiro,
    humano, material, sistema de alerta ou treinamento).

`ibge_malha_rs`
    Malha municipal do RS em GeoJSON (API de malhas v3, qualidade minima).
    Sem ela nao existe mapa municipal — so lista.

`ibge_pop_rs`
    Populacao estimada por municipio (agregado 6579). Vem tambem dentro do
    MUNIC (coluna PopMun); ingerida a parte porque o MUNIC congela em 2024 e
    a populacao e atualizada todo ano.

Uso:
    python -m src.ingest.ibge_rs            # todas
    python -m src.ingest.ibge_rs malha      # uma
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import sys
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[2]
DATA_RAW = REPO_ROOT / "data" / "raw"
DATA_INTERIM = REPO_ROOT / "data" / "interim"

CACHE_HOURS = 24 * 30  # MUNIC e anual; malha muda quando ha mudanca territorial
TIMEOUT_S = 300  # o xlsx do MUNIC tem ~25 MB
USER_AGENT = "climate-rs-engine/1"


class IngestError(RuntimeError):
    """Erro de ingestao — a mensagem sempre nomeia fonte + etapa."""


@dataclass(frozen=True)
class Source:
    source_id: str
    url: str
    file_ext: str
    notes: str


MUNIC = Source(
    source_id="ibge_munic_rs",
    url="https://ftp.ibge.gov.br/Perfil_Municipios/2024/Base_de_Dados/Base_MUNIC_2024_20251107.xlsx",
    file_ext="xlsx",
    notes=(
        "IBGE MUNIC 2024, aba 'Evento climático RS' (suplemento sobre o evento "
        "de tempestade iniciado em 26/04/2024). Resposta declarada pela "
        "prefeitura de cada municipio."
    ),
)

MALHA = Source(
    source_id="ibge_malha_rs",
    url=(
        "https://servicodados.ibge.gov.br/api/v3/malhas/estados/43"
        "?formato=application/vnd.geo+json&intrarregiao=municipio&qualidade=minima"
    ),
    file_ext="geojson",
    notes="Malha municipal do RS, API de malhas v3, qualidade minima (leve o bastante para o front).",
)

POP = Source(
    source_id="ibge_pop_rs",
    url=(
        "https://servicodados.ibge.gov.br/api/v3/agregados/6579/periodos/2024"
        "/variaveis/9324?localidades=N6[N3[43]]"
    ),
    file_ext="json",
    notes="Populacao residente estimada por municipio, agregado SIDRA 6579, ano 2024.",
)


# ---------------------------------------------------------------------------
# Infra: HTTP, cache por janela, gravacao do bruto com proveniencia
# ---------------------------------------------------------------------------
def _http_get(source: Source) -> bytes:
    if os.environ.get("CLIMATE_OFFLINE") == "1":
        raise IngestError(f"{source.source_id}: fetch bloqueado (CLIMATE_OFFLINE=1)")
    req = urllib.request.Request(source.url, headers={"User-Agent": USER_AGENT})
    try:
        with urllib.request.urlopen(req, timeout=TIMEOUT_S) as resp:
            body = resp.read()
            # A API de malhas do IBGE responde gzip mesmo sem Accept-Encoding.
            # Descomprimir ANTES de gravar o bruto: o payload em disco tem de
            # ser legivel por qualquer ferramenta, nao so pelo nosso parser.
            if body[:2] == b"\x1f\x8b":
                import gzip

                body = gzip.decompress(body)
            return body
    except (urllib.error.URLError, urllib.error.HTTPError, TimeoutError, OSError) as exc:
        raise IngestError(f"{source.source_id}: falha na etapa fetch, url={source.url}: {exc}") from exc


def _latest_raw_dir(source_id: str) -> Path | None:
    base = DATA_RAW / source_id
    if not base.exists():
        return None
    dirs = sorted((p for p in base.iterdir() if p.is_dir()), reverse=True)
    return dirs[0] if dirs else None


def _payload(source: Source) -> tuple[bytes, Path]:
    """Bruto respeitando cache e modo offline. Retorna (bytes, diretorio)."""
    offline = os.environ.get("CLIMATE_OFFLINE") == "1"
    latest = _latest_raw_dir(source.source_id)

    if latest is not None:
        found = next(latest.glob("payload.*"), None)
        prov_path = latest / "provenance.json"
        if found is not None and prov_path.exists():
            if offline:
                return found.read_bytes(), latest
            prov = json.loads(prov_path.read_text(encoding="utf-8"))
            age_h = (pd.Timestamp.now(tz="UTC") - pd.Timestamp(prov["fetched_at"])).total_seconds() / 3600
            if age_h <= CACHE_HOURS:
                return found.read_bytes(), latest

    if offline:
        raise IngestError(
            f"{source.source_id}: modo offline sem download previo em "
            f"{DATA_RAW / source.source_id} — impossivel prosseguir"
        )

    content = _http_get(source)
    out_dir = DATA_RAW / source.source_id / pd.Timestamp.now(tz="UTC").strftime("%Y%m%dT%H%M%S")
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / f"payload.{source.file_ext}").write_bytes(content)
    (out_dir / "provenance.json").write_text(
        json.dumps(
            {
                "url": source.url,
                "fetched_at": pd.Timestamp.now(tz="UTC").isoformat(),
                "sha256": hashlib.sha256(content).hexdigest(),
                "product_version": "2024",
                "ingestor_version": INGESTOR_VERSION,
                "rows": None,  # preenchido pelo parse via _stamp_rows
                "notes": source.notes,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return content, out_dir


INGESTOR_VERSION = "1"


def _stamp_rows(raw_dir: Path, rows: int) -> None:
    """Grava a contagem de linhas na proveniencia depois do parse.

    O bruto e gravado antes de saber quantas linhas ele rende (parse pode
    falhar). Voltar aqui mantem /health/sources util sem adiar a gravacao
    do payload, que e o artefato que nao pode se perder.
    """
    prov_path = raw_dir / "provenance.json"
    if not prov_path.exists():
        return
    prov = json.loads(prov_path.read_text(encoding="utf-8"))
    prov["rows"] = rows
    prov_path.write_text(json.dumps(prov, indent=2, ensure_ascii=False), encoding="utf-8")


# ---------------------------------------------------------------------------
# MUNIC 2024 — aba "Evento climático RS"
# ---------------------------------------------------------------------------
SHEET_EVENTO = "Evento climático RS"

# Mapa codigo MUNIC -> nome estavel de coluna. Escrito por extenso de
# proposito: `Mers099` nao diz nada seis meses depois, e o custo de errar a
# leitura de uma variavel aqui e um numero errado numa central de risco.
# Fonte dos rotulos: aba 'Dicionário' do proprio MUNIC 2024.
MUNIC_COLS: dict[str, str] = {
    "CodMun": "cod_mun",
    "Desc Mun": "municipio",
    "PopMun": "populacao",
    "Regiao": "regiao_ibge",
    # 1. Sistema de alerta e plano de contingencia
    "Mers01": "atingido",                    # foi atingido pelo evento
    "Mers02": "alerta_emitido",              # autoridades emitiram alertas
    "Mers031": "alerta_sms",
    "Mers032": "alerta_app",
    "Mers033": "alerta_sirene",
    "Mers034": "alerta_outro",
    "Mers05": "alerta_alcance",              # faixa da populacao alcancada
    "Mers06": "plano_contingencia",          # existencia
    "Mers07": "plano_executado",
    "Mers082": "falta_recurso_financeiro",   # motivos de nao execucao
    "Mers083": "falta_recurso_humano",
    "Mers084": "falta_recurso_material",
    "Mers085": "falta_sistema_alerta",
    "Mers086": "falta_treinamento",
    # 2. Dimensionamento do impacto — perigos ocorridos
    "Mers091": "oc_deslizamento",
    "Mers092": "oc_corrida_massa",
    "Mers093": "oc_queda_blocos",
    "Mers094": "oc_queda_barreira",
    "Mers095": "oc_desabamento",
    "Mers096": "oc_erosao",
    "Mers097": "oc_solapamento_margem",
    "Mers098": "oc_inundacao",
    "Mers099": "oc_enchente_enxurrada",
    "Mers0910": "oc_alagamento",
    # grupos expostos
    "Mers1110": "exp_favelas",
    "Mers116": "exp_situacao_rua",
    "Mers117": "exp_comunidades_tradicionais",
    "Mers118": "exp_deficiencia",
    "Mers111": "exp_criancas",
    "Mers112": "exp_mulheres",
    "Mers113": "exp_lgbtqia",
    "Mers114": "exp_gestantes",
    "Mers115": "exp_populacao_negra",
    "Mers119": "exp_doencas_cronicas",
    # 2b. Impacto sobre o sistema de saude
    "Mers121": "saude_estruturas_afetadas",
    "Mers122": "saude_campanha_vacinal",
    "Mers123": "saude_danos_equipamentos",
    "Mers124": "saude_atendimento_suspenso",
    "Mers125": "saude_remanejamento_pacientes",
    "Mers126": "saude_combustivel",
    "Mers135": "saude_doencas_inundacao",
    "Mers1311": "dano_acolhimento_institucional",
    # 3. Resposta prestada — o que o municipio conseguiu entregar
    "Mers161": "resp_equipes_resgate",
    "Mers162": "resp_transporte_vitimas",
    "Mers163": "resp_ambulancias",
    "Mers164": "resp_recursos_sus",
    "Mers165": "resp_abrigo",
    "Mers166": "resp_apoio_psicologico",
    "Mers167": "resp_alimentos_agua",
    "Mers169": "resp_farmacos",
    "Mers1613": "resp_medicamentos_atencao",
    "Mers1616": "resp_visita_domiciliar",
    "Mers1617": "resp_limpeza_vias",
    # 3b. Autonomia logistica — as sete perguntas ordinais (nao sao Sim/Nao)
    "Mers17": "log_tempo_primeira_resposta",
    "Mers18": "log_recursos_durante",
    "Mers19": "log_dias_fornecimento",
    "Mers20": "log_cobertura_bairros",
    "Mers21": "log_72h_criticas",
    "Mers22": "log_alimentacao",
    "Mers23": "log_cooperacao",
    # 3. Danos
    "Mers131": "dano_obitos",
    "Mers132": "dano_desaparecidos",
    "Mers133": "dano_desabrigados",
    "Mers134": "dano_feridos",
    "Mers136": "dano_portos_aeroportos",
    "Mers137": "dano_barragens",
    "Mers138": "dano_viario",
    "Mers139": "dano_industria",
    "Mers1310": "dano_agropecuaria",
    "Mers1313": "areas_ilhadas",
    "Mers1314": "risco_contaminacao_quimica",
}

# Colunas cujo valor no MUNIC e Sim/Não/- e viram booleano com ausencia
# preservada (None != False: "-" significa pergunta nao aplicavel porque o
# municipio nao foi atingido, nao "respondeu que nao").
_SIM = {"sim"}
_NAO = {"não", "nao"}


def _tri(value: object) -> bool | None:
    """Sim -> True, Não -> False, qualquer outra coisa -> None.

    O "qualquer outra coisa" agrupa quatro respostas que o MUNIC distingue:
    `-` (nao aplicavel, porque o municipio nao foi atingido), `Nao sabe
    informar`, `Nao informou` e `Recusa`. Todas viram None de proposito: do
    ponto de vista do indice, nenhuma delas e um "nao" — e tratar "nao sabe"
    como ausencia de problema seria a leitura mais otimista possivel do
    silencio, exatamente o erro que a cobertura minima do modelo evita.

    A distincao entre elas se perde aqui e esta declarada como limite: quem
    precisar separar "nao aplicavel" de "nao sabe" tem o bruto em disco.
    """
    if value is None:
        return None
    s = str(value).strip().lower()
    if s in _SIM:
        return True
    if s in _NAO:
        return False
    return None


def parse_munic(raw: bytes) -> pd.DataFrame:
    import openpyxl  # dependencia so deste parser — nao carregar no import do modulo

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if SHEET_EVENTO not in wb.sheetnames:
        raise IngestError(
            f"{MUNIC.source_id}: aba '{SHEET_EVENTO}' ausente do xlsx "
            f"(abas: {wb.sheetnames}) — o IBGE mudou a estrutura da base"
        )
    ws = wb[SHEET_EVENTO]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    df = pd.DataFrame(list(rows), columns=header)

    missing = [c for c in MUNIC_COLS if c not in df.columns]
    if missing:
        raise IngestError(
            f"{MUNIC.source_id}: colunas ausentes na aba '{SHEET_EVENTO}': {missing}"
        )

    out = df[list(MUNIC_COLS)].rename(columns=MUNIC_COLS).copy()
    out["cod_mun"] = out["cod_mun"].astype(int)
    out["populacao"] = pd.to_numeric(out["populacao"], errors="coerce").astype("Int64")
    out["municipio"] = out["municipio"].astype(str).str.strip()

    # Colunas ORDINAIS, nao booleanas: alerta_alcance e faixa textual ("Menos
    # de 5%"), e as sete `log_*` sao escalas proprias do MUNIC ("Igual a
    # procura", "Menor em 15%", "Nao houve/nao necessitou"...). Ficam como
    # texto — traduzir escala ordinal para numero e decisao do MODELO, nao da
    # ingestao: a Camada 1 registra, nao interpreta.
    ORDINAIS = {"alerta_alcance", *[c for c in out.columns if c.startswith("log_")]}
    for col in ORDINAIS:
        out[col] = out[col].astype(str).str.strip()

    nao_bool = {"cod_mun", "municipio", "populacao", "regiao_ibge", *ORDINAIS}
    bool_cols = [c for c in out.columns if c not in nao_bool]
    for col in bool_cols:
        out[col] = out[col].map(_tri).astype("boolean")

    if len(out) != 497:
        raise IngestError(
            f"{MUNIC.source_id}: esperado 497 municipios do RS, obtido {len(out)}"
        )
    return out.sort_values("cod_mun").reset_index(drop=True)


# ---------------------------------------------------------------------------
# Malha municipal e populacao
# ---------------------------------------------------------------------------
def parse_malha(raw: bytes) -> dict:
    geo = json.loads(raw.decode("utf-8"))
    feats = geo.get("features", [])
    if len(feats) != 497:
        raise IngestError(f"{MALHA.source_id}: esperado 497 poligonos, obtido {len(feats)}")
    return geo


def parse_pop(raw: bytes) -> pd.DataFrame:
    payload = json.loads(raw.decode("utf-8"))
    try:
        series = payload[0]["resultados"][0]["series"]
    except (IndexError, KeyError) as exc:
        raise IngestError(f"{POP.source_id}: estrutura inesperada do agregado 6579: {exc}") from exc
    recs = []
    for s in series:
        loc = s["localidade"]
        valores = s["serie"]
        raw_v = next(iter(valores.values()), None)
        try:
            pop = int(raw_v)
        except (TypeError, ValueError):
            pop = None
        recs.append({"cod_mun": int(loc["id"]), "municipio": loc["nome"], "populacao": pop})
    return pd.DataFrame(recs).sort_values("cod_mun").reset_index(drop=True)


# ---------------------------------------------------------------------------
# Orquestracao
# ---------------------------------------------------------------------------
def run_munic() -> pd.DataFrame:
    raw, raw_dir = _payload(MUNIC)
    try:
        df = parse_munic(raw)
    except IngestError:
        raise
    except Exception as exc:
        raise IngestError(f"{MUNIC.source_id}: falha na etapa parse ({raw_dir}): {exc}") from exc
    DATA_INTERIM.mkdir(parents=True, exist_ok=True)
    df.to_parquet(DATA_INTERIM / f"{MUNIC.source_id}.parquet", index=False)
    _stamp_rows(raw_dir, len(df))
    return df


def run_malha() -> dict:
    raw, raw_dir = _payload(MALHA)
    geo = parse_malha(raw)
    DATA_INTERIM.mkdir(parents=True, exist_ok=True)
    (DATA_INTERIM / f"{MALHA.source_id}.geojson").write_text(
        json.dumps(geo, ensure_ascii=False), encoding="utf-8"
    )
    _stamp_rows(raw_dir, len(geo["features"]))
    return geo


def run_pop() -> pd.DataFrame:
    raw, raw_dir = _payload(POP)
    df = parse_pop(raw)
    DATA_INTERIM.mkdir(parents=True, exist_ok=True)
    df.to_parquet(DATA_INTERIM / f"{POP.source_id}.parquet", index=False)
    _stamp_rows(raw_dir, len(df))
    return df


RUNNERS = {"munic": run_munic, "malha": run_malha, "pop": run_pop}


def main(argv: list[str]) -> int:
    targets = argv[1:] or list(RUNNERS)
    unknown = [t for t in targets if t not in RUNNERS]
    if unknown:
        print(f"alvo desconhecido: {unknown}; disponiveis: {list(RUNNERS)}", file=sys.stderr)
        return 2
    for name in targets:
        try:
            result = RUNNERS[name]()
        except IngestError as exc:
            print(f"[FALHA] {name}: {exc}", file=sys.stderr)
            return 1
        n = len(result["features"]) if isinstance(result, dict) else len(result)
        print(f"[OK] {name}: {n} registros")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
